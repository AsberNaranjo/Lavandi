"""Construye el listado de lavanderías y tintorerías del territorio del contrato.

Fuente única: el Directorio Estadístico Nacional de Unidades Económicas (DENUE) del INEGI,
actividad 812210 "Lavanderías y tintorerías". Descarga los archivos de Durango (10) y
Coahuila (05), se queda con todo el estado de Durango y con Torreón y Matamoros, y escribe
un Excel para el cliente con un listado listo para llamar y un resumen por plaza.

Uso:
    python -I construir_listado.py [carpeta_de_descarga] [archivo_de_salida.xlsx]

Si la carpeta ya trae los .zip del INEGI (denue_10.zip y denue_05.zip), no se vuelven a bajar.
"""
import csv
import io
import re
import sys
import unicodedata
import urllib.request
import zipfile
from collections import Counter
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

URL_DENUE = "https://www.inegi.org.mx/contenidos/masiva/denue/denue_{ent}_csv.zip"
ACTIVIDAD = "812210"
MUNICIPIOS_COAHUILA = {"torreon", "matamoros"}

# Orden en que se presentan las plazas.
PLAZAS = [
    "Durango capital",
    "Gómez Palacio y Lerdo",
    "Torreón",
    "Matamoros",
    "Resto del estado de Durango",
]

SIN_TIPO = "No se sabe por el nombre"

RAIZ = Path(__file__).resolve().parents[2]
SALIDA_POR_OMISION = RAIZ / "vault" / "negocio" / "listado-de-lavanderias.xlsx"

FUENTE = "Arial"
COLOR_TITULO = "1F4E5A"
COLOR_ENCABEZADO = "DCE8EA"
COLOR_LLENAR = "FFF2B3"
BORDE = Border(bottom=Side(style="thin", color="C9D3D6"))


def normalizar(texto):
    texto = unicodedata.normalize("NFKD", texto or "")
    return "".join(c for c in texto if not unicodedata.combining(c)).lower().strip()


def limpio(valor):
    valor = (valor or "").strip()
    return "" if valor.upper() in {"", "SN", "S/N", "NINGUNO", "NINGUNA", "0"} else valor


def bajar(carpeta, ent):
    destino = carpeta / f"denue_{ent}.zip"
    if not destino.exists():
        print(f"Descargando {URL_DENUE.format(ent=ent)} …")
        urllib.request.urlretrieve(URL_DENUE.format(ent=ent), destino)
    return destino


def leer(zip_path):
    with zipfile.ZipFile(zip_path) as z:
        nombre = next(n for n in z.namelist() if re.search(r"conjunto_de_datos/denue_inegi_.*\.csv$", n))
        metadatos = next((n for n in z.namelist() if n.endswith("metadatos_denue.txt")), None)
        edicion = ""
        if metadatos:
            texto = z.read(metadatos).decode("utf-8-sig", errors="replace")
            m = re.search(r"^Title:.*?(\d{2}_\d{4})", texto, re.M)
            edicion = m.group(1) if m else ""
        with z.open(nombre) as f:
            filas = list(csv.DictReader(io.TextIOWrapper(f, encoding="latin1", newline="")))
    return filas, edicion


def plaza_de(fila):
    mun = normalizar(fila["municipio"])
    if fila["cve_ent"] == "05":
        return "Torreón" if mun == "torreon" else "Matamoros"
    if mun == "durango":
        return "Durango capital"
    if mun in {"gomez palacio", "lerdo"}:
        return "Gómez Palacio y Lerdo"
    return "Resto del estado de Durango"


def tipo_por_nombre(fila):
    """Deduce del nombre qué servicios anuncia el negocio. No es una confirmación."""
    nombre = normalizar(f"{fila['nom_estab']} {fila['raz_social']}")
    servicios = []
    if re.search(r"lava|laundr|wash", nombre):
        servicios.append("Lavandería")
    if re.search(r"tintor|dry ?clean|5 ?asec", nombre):
        servicios.append("Tintorería")
    if re.search(r"planch|\biron\b", nombre):
        servicios.append("Planchaduría")
    if not servicios:
        return SIN_TIPO
    primero, *resto = servicios
    if not resto:
        return primero
    resto = [s.lower() for s in resto]
    return ", ".join([primero] + resto[:-1]) + " y " + resto[-1]


def telefono(valor):
    digitos = re.sub(r"\D", "", valor or "")
    if len(digitos) == 10:
        return f"{digitos[:3]} {digitos[3:6]} {digitos[6:]}"
    return digitos


def direccion(fila):
    calle = " ".join(x for x in (limpio(fila["tipo_vial"]).title(), limpio(fila["nom_vial"]).title()) if x)
    numero = " ".join(x for x in (limpio(fila["numero_ext"]), limpio(fila["letra_ext"])) if x)
    interior = " ".join(x for x in (limpio(fila["numero_int"]), limpio(fila["letra_int"])) if x)
    partes = [calle + (f" {numero}" if numero else " sin número")]
    if interior:
        partes.append(f"interior {interior}")
    if limpio(fila["nom_CenCom"]):
        partes.append(limpio(fila["nom_CenCom"]).title())
    return ", ".join(partes)


def colonia(fila):
    return " ".join(x for x in (limpio(fila["tipo_asent"]).title(), limpio(fila["nomb_asent"]).title()) if x)


def registros(carpeta):
    lavanderias, otras, ediciones = [], [], set()
    for ent in ("10", "05"):
        filas, edicion = leer(bajar(carpeta, ent))
        ediciones.add(edicion)
        for fila in filas:
            if ent == "05" and normalizar(fila["municipio"]) not in MUNICIPIOS_COAHUILA:
                continue
            nombre = normalizar(f"{fila['nom_estab']} {fila['raz_social']}")
            if fila["codigo_act"] == ACTIVIDAD:
                lavanderias.append(fila)
            elif re.search(r"lavander|tintorer", nombre):
                otras.append(fila)
    return lavanderias, otras, ediciones


def estilo_encabezado(celda):
    celda.font = Font(name=FUENTE, bold=True, color=COLOR_TITULO)
    celda.fill = PatternFill("solid", fgColor=COLOR_ENCABEZADO)
    celda.alignment = Alignment(wrap_text=True, vertical="center")


def titulo(hoja, texto, subtitulo=None):
    hoja["A1"] = texto
    hoja["A1"].font = Font(name=FUENTE, bold=True, size=16, color=COLOR_TITULO)
    if subtitulo:
        hoja["A2"] = subtitulo
        hoja["A2"].font = Font(name=FUENTE, italic=True, size=10, color="5A6B70")


def fila_listado(i, fila):
    lat, lon = fila["latitud"], fila["longitud"]
    return [
        i,
        plaza_de(fila),
        fila["municipio"].strip(),
        fila["localidad"].strip(),
        fila["nom_estab"].strip(),
        limpio(fila["raz_social"]),
        tipo_por_nombre(fila),
        fila["per_ocu"].strip(),
        direccion(fila),
        colonia(fila),
        limpio(fila["cod_postal"]),
        telefono(fila["telefono"]),
        limpio(fila["correoelec"]).lower(),
        limpio(fila["www"]).lower(),
        f"https://www.google.com/maps/search/?api=1&query={lat},{lon}",
        float(lat),
        float(lon),
        fila["id"],
        "",
        "",
        "",
    ]


COLUMNAS = [
    ("N.º", 6), ("Plaza", 22), ("Municipio", 18), ("Localidad", 20), ("Nombre del negocio", 34),
    ("Razón social", 28), ("Qué es, según su nombre", 22), ("Personas que trabajan ahí", 14),
    ("Calle y número", 36), ("Colonia", 28), ("Código postal", 10), ("Teléfono", 15),
    ("Correo", 28), ("Página web o red social", 28), ("Ver en el mapa", 14), ("Latitud", 11),
    ("Longitud", 11), ("Clave del INEGI", 12),
    ("¿Ya se le llamó?", 12), ("Resultado de la llamada", 22), ("Notas", 40),
]
COL = {nombre: get_column_letter(i + 1) for i, (nombre, _) in enumerate(COLUMNAS)}
RESULTADOS = ["Le interesa", "Quiere más información", "No le interesa", "No contestó",
              "Número equivocado", "Ya cerró"]


def hoja_listado(wb, lavanderias):
    hoja = wb.create_sheet("Listado")
    orden = {p: i for i, p in enumerate(PLAZAS)}
    lavanderias.sort(key=lambda f: (orden[plaza_de(f)], normalizar(f["municipio"]), normalizar(f["nom_estab"])))
    hoja.append([c for c, _ in COLUMNAS])
    for i, (_, ancho) in enumerate(COLUMNAS, start=1):
        estilo_encabezado(hoja.cell(row=1, column=i))
        hoja.column_dimensions[get_column_letter(i)].width = ancho
    for i, fila in enumerate(lavanderias, start=1):
        hoja.append(fila_listado(i, fila))
    ultima = hoja.max_row
    for r in range(2, ultima + 1):
        for c in range(1, len(COLUMNAS) + 1):
            hoja.cell(row=r, column=c).font = Font(name=FUENTE, size=10)
        enlace = hoja[f"{COL['Ver en el mapa']}{r}"]
        enlace.hyperlink = enlace.value
        enlace.value = "Abrir mapa"
        enlace.font = Font(name=FUENTE, size=10, color="1F5FAF", underline="single")
        for nombre in ("¿Ya se le llamó?", "Resultado de la llamada", "Notas"):
            hoja[f"{COL[nombre]}{r}"].fill = PatternFill("solid", fgColor=COLOR_LLENAR)
        hoja[f"{COL['Latitud']}{r}"].number_format = "0.000000"
        hoja[f"{COL['Longitud']}{r}"].number_format = "0.000000"
    hoja.row_dimensions[1].height = 32
    hoja.freeze_panes = "F2"
    hoja.auto_filter.ref = f"A1:{get_column_letter(len(COLUMNAS))}{ultima}"

    si_no = DataValidation(type="list", formula1='"Sí,No"', allow_blank=True)
    resultado = DataValidation(type="list", formula1='"' + ",".join(RESULTADOS) + '"', allow_blank=True)
    hoja.add_data_validation(si_no)
    hoja.add_data_validation(resultado)
    si_no.add(f"{COL['¿Ya se le llamó?']}2:{COL['¿Ya se le llamó?']}{ultima}")
    resultado.add(f"{COL['Resultado de la llamada']}2:{COL['Resultado de la llamada']}{ultima}")
    return ultima


def rango(nombre, ultima):
    return f"Listado!${COL[nombre]}$2:${COL[nombre]}${ultima}"


def hoja_resumen(wb, ultima):
    hoja = wb.create_sheet("Resumen", 1)
    titulo(hoja, "Resumen por plaza", "Las cifras se recalculan solas a partir de la hoja Listado.")
    encabezados = ["Plaza", "Negocios", "Con teléfono", "Con correo", "Con página o red social",
                   "Su nombre dice lavandería", "Su nombre dice tintorería", "Su nombre dice planchaduría",
                   SIN_TIPO, "Ya se les llamó", "Les interesa"]
    fila0 = 4
    for c, texto in enumerate(encabezados, start=1):
        celda = hoja.cell(row=fila0, column=c, value=texto)
        estilo_encabezado(celda)
        hoja.column_dimensions[get_column_letter(c)].width = 26 if c == 1 else 14
    plaza = rango("Plaza", ultima)
    for k, nombre in enumerate(PLAZAS, start=1):
        r = fila0 + k
        a = f"$A{r}"
        formulas = [
            nombre,
            f"=COUNTIF({plaza},{a})",
            f'=COUNTIFS({plaza},{a},{rango("Teléfono", ultima)},"<>")',
            f'=COUNTIFS({plaza},{a},{rango("Correo", ultima)},"<>")',
            f'=COUNTIFS({plaza},{a},{rango("Página web o red social", ultima)},"<>")',
            f'=COUNTIFS({plaza},{a},{rango("Qué es, según su nombre", ultima)},"*avandería*")',
            f'=COUNTIFS({plaza},{a},{rango("Qué es, según su nombre", ultima)},"*intorería*")',
            f'=COUNTIFS({plaza},{a},{rango("Qué es, según su nombre", ultima)},"*lanchaduría*")',
            f'=COUNTIFS({plaza},{a},{rango("Qué es, según su nombre", ultima)},"{SIN_TIPO}")',
            f'=COUNTIFS({plaza},{a},{rango("¿Ya se le llamó?", ultima)},"Sí")',
            f'=COUNTIFS({plaza},{a},{rango("Resultado de la llamada", ultima)},"Le interesa")',
        ]
        for c, valor in enumerate(formulas, start=1):
            celda = hoja.cell(row=r, column=c, value=valor)
            celda.font = Font(name=FUENTE, size=11)
            celda.border = BORDE
    total = fila0 + len(PLAZAS) + 1
    hoja.cell(row=total, column=1, value="Todo el territorio")
    for c in range(2, len(encabezados) + 1):
        letra = get_column_letter(c)
        hoja.cell(row=total, column=c, value=f"=SUM({letra}{fila0 + 1}:{letra}{total - 1})")
    for c in range(1, len(encabezados) + 1):
        hoja.cell(row=total, column=c).font = Font(name=FUENTE, bold=True, size=11)
    hoja.row_dimensions[fila0].height = 32

    nota = total + 2
    textos = [
        "Cómo leer este resumen",
        "• «Negocios» cuenta cada local por separado: si un negocio tiene tres sucursales, aparecen tres.",
        "• Las columnas «Su nombre dice…» salen solo del nombre del negocio. Un negocio puede contar en más de una (por ejemplo, «Lavandería y planchaduría»), y uno que solo dice «Lavandería» puede también planchar o lavar en seco: se confirma en la llamada.",
        "• «Ya se les llamó» y «Les interesa» se llenan solas cuando alguien marca esas columnas en la hoja Listado.",
    ]
    for i, texto in enumerate(textos):
        celda = hoja.cell(row=nota + i, column=1, value=texto)
        celda.font = Font(name=FUENTE, bold=(i == 0), size=10)
    hoja.freeze_panes = f"A{fila0 + 1}"


def hoja_municipios(wb, lavanderias, ultima):
    hoja = wb.create_sheet("Por municipio", 2)
    titulo(hoja, "Negocios por municipio", "Ordenados de mayor a menor.")
    for c, texto in enumerate(["Plaza", "Municipio", "Negocios", "Con teléfono"], start=1):
        estilo_encabezado(hoja.cell(row=4, column=c, value=texto))
    hoja.column_dimensions["A"].width = 28
    hoja.column_dimensions["B"].width = 26
    hoja.column_dimensions["C"].width = 12
    hoja.column_dimensions["D"].width = 14
    conteo = Counter((plaza_de(f), f["municipio"].strip()) for f in lavanderias)
    orden = {p: i for i, p in enumerate(PLAZAS)}
    pares = sorted(conteo, key=lambda k: (orden[k[0]], -conteo[k], k[1]))
    mun = rango("Municipio", ultima)
    for i, (plaza, municipio) in enumerate(pares, start=5):
        hoja.cell(row=i, column=1, value=plaza)
        hoja.cell(row=i, column=2, value=municipio)
        hoja.cell(row=i, column=3, value=f"=COUNTIF({mun},$B{i})")
        hoja.cell(row=i, column=4, value=f'=COUNTIFS({mun},$B{i},{rango("Teléfono", ultima)},"<>")')
        for c in range(1, 5):
            hoja.cell(row=i, column=c).font = Font(name=FUENTE, size=10)
            hoja.cell(row=i, column=c).border = BORDE
    hoja.freeze_panes = "A5"


def hoja_otras(wb, otras):
    hoja = wb.create_sheet("Otras posibles")
    titulo(hoja, "Otros negocios que podrían lavar ropa",
           "El INEGI los registró con otra actividad, pero su nombre dice lavandería o tintorería. No están en el Listado ni en el Resumen.")
    encabezados = ["Municipio", "Nombre del negocio", "Actividad registrada", "Calle y número", "Colonia",
                   "Teléfono", "Clave del INEGI"]
    anchos = [18, 36, 40, 36, 28, 15, 12]
    for c, (texto, ancho) in enumerate(zip(encabezados, anchos), start=1):
        estilo_encabezado(hoja.cell(row=4, column=c, value=texto))
        hoja.column_dimensions[get_column_letter(c)].width = ancho
    for r, fila in enumerate(sorted(otras, key=lambda f: (normalizar(f["municipio"]), normalizar(f["nom_estab"]))), start=5):
        valores = [fila["municipio"].strip(), fila["nom_estab"].strip(), fila["nombre_act"].strip(),
                   direccion(fila), colonia(fila), telefono(fila["telefono"]), fila["id"]]
        for c, v in enumerate(valores, start=1):
            hoja.cell(row=r, column=c, value=v).font = Font(name=FUENTE, size=10)


def hoja_leeme(wb, total, sin_telefono):
    hoja = wb.active
    hoja.title = "Léeme"
    titulo(hoja, "Lavanderías y tintorerías del territorio",
           "Todo el estado de Durango, más Torreón y Matamoros, Coahuila. Preparado por Nazmelia.")
    hoja.column_dimensions["A"].width = 110
    bloques = [
        ("Qué es este archivo", [
            f"La lista de los {total} negocios de lavandería y tintorería que el INEGI tiene registrados en el territorio, con su dirección, teléfono, correo y página cuando el INEGI los tiene.",
            "Sirve para buscar lavanderías socias: a quién llamar, en qué zona y cómo va cada llamada.",
        ]),
        ("De dónde salen los datos", [
            "Del Directorio Estadístico Nacional de Unidades Económicas del INEGI, en su edición más reciente, actividad «Lavanderías y tintorerías».",
            "Es un registro público y oficial. Los datos los capturó el INEGI; nosotros solo los ordenamos y los dejamos listos para usar.",
        ]),
        ("Qué hay en cada hoja", [
            "Resumen: cuántos negocios hay en cada plaza y cuántos tienen teléfono, correo o página. También lleva la cuenta de las llamadas.",
            "Por municipio: lo mismo, municipio por municipio.",
            "Listado: un negocio por renglón. Se puede filtrar por plaza, municipio o colonia con las flechas del encabezado.",
            "Otras posibles: negocios que el INEGI registró con otra actividad pero cuyo nombre dice lavandería o tintorería. Vale la pena revisarlos aparte.",
        ]),
        ("Cómo registrar las llamadas", [
            "En la hoja Listado, las tres columnas amarillas del final son para llenarse a mano:",
            "   «¿Ya se le llamó?»: elegir Sí o No.",
            "   «Resultado de la llamada»: elegir una opción de la lista (Le interesa, Quiere más información, No le interesa, No contestó, Número equivocado, Ya cerró).",
            "   «Notas»: texto libre. Ejemplo: «Habló con la dueña, Sra. López. Lava por kilo, ya entrega a domicilio. Volver a llamar el lunes».",
            "El Resumen se actualiza solo con lo que se marque ahí.",
        ]),
        ("Cuidado con estos datos", [
            "El registro del INEGI puede estar desfasado: hay negocios que ya cerraron, que cambiaron de teléfono o que abrieron después y todavía no aparecen.",
            f"{sin_telefono} de {total} no tienen teléfono registrado. Para esos, el enlace «Abrir mapa» lleva a su ubicación y ahí a veces aparece un teléfono más reciente.",
            "«Qué es, según su nombre» es una deducción a partir del nombre, no una confirmación.",
            "Este archivo no mide si los negocios tienen presencia en internet ni cuántas personas buscan el servicio. Eso se trabaja aparte.",
        ]),
    ]
    r = 4
    for encabezado, lineas in bloques:
        hoja.cell(row=r, column=1, value=encabezado).font = Font(name=FUENTE, bold=True, size=12, color=COLOR_TITULO)
        r += 1
        for linea in lineas:
            celda = hoja.cell(row=r, column=1, value=linea)
            celda.font = Font(name=FUENTE, size=11)
            celda.alignment = Alignment(wrap_text=True, vertical="top")
            r += 1
        r += 1


def main():
    carpeta = Path(sys.argv[1]) if len(sys.argv) > 1 else Path.cwd() / "descarga_denue"
    salida = Path(sys.argv[2]) if len(sys.argv) > 2 else SALIDA_POR_OMISION
    carpeta.mkdir(parents=True, exist_ok=True)
    lavanderias, otras, ediciones = registros(carpeta)

    wb = Workbook()
    hoja_leeme(wb, len(lavanderias), sum(1 for f in lavanderias if not telefono(f["telefono"])))
    ultima = hoja_listado(wb, lavanderias)
    hoja_resumen(wb, ultima)
    hoja_municipios(wb, lavanderias, ultima)
    hoja_otras(wb, otras)
    wb.properties.creator = "Nazmelia"
    wb.properties.title = "Lavanderías y tintorerías del territorio"
    salida.parent.mkdir(parents=True, exist_ok=True)
    wb.save(salida)

    por_plaza = Counter(plaza_de(f) for f in lavanderias)
    print(f"Edición del DENUE: {', '.join(sorted(ediciones))}")
    print(f"Negocios (actividad {ACTIVIDAD}): {len(lavanderias)}")
    for p in PLAZAS:
        print(f"  {p}: {por_plaza[p]}")
    print(f"Con teléfono: {sum(1 for f in lavanderias if telefono(f['telefono']))}")
    print(f"Con correo: {sum(1 for f in lavanderias if limpio(f['correoelec']))}")
    print(f"Con página: {sum(1 for f in lavanderias if limpio(f['www']))}")
    print(f"Otras posibles: {len(otras)}")
    print(f"Tipos: {Counter(tipo_por_nombre(f) for f in lavanderias)}")
    print(f"Escrito: {salida}")


if __name__ == "__main__":
    main()
