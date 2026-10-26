"""Versión para EL CLIENTE del listado de lavanderías: solo lo que sirve para prospectar.

Usa los mismos datos y reglas que construir_listado.py, pero deja fuera todo lo técnico
(coordenadas, claves del INEGI, razón social, clasificación por nombre) y separa a quién se le
puede llamar de quién hay que visitar. Lenguaje llano y sin fechas (C-108, C-109).

Uso:
    python -I version_cliente.py [carpeta_de_descarga] [archivo_de_salida.xlsx]
"""
import re
import sys
from collections import Counter
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

sys.path.insert(0, str(Path(__file__).resolve().parent))  # con `python -I` la carpeta no entra sola

from construir_listado import (PLAZAS, RAIZ, TIPOS, TIPOS_INFO, colonia, direccion, limpio, normalizar, plaza_de,
                               registros, telefono, tipos_de_prospecto)

SALIDA_POR_OMISION = RAIZ / "vault" / "cliente" / "comunicaciones" / "lavanderias-para-invitar.xlsx"

# Colores de Lavandi (vault/design-system/tokens.json).
AZUL = "297899"
AZUL_CLARO = "E3F1F6"
AMARILLO = "FFF4CC"
GRIS = "515A64"
FUENTE = "Arial"
LINEA = Border(bottom=Side(style="thin", color="DCE0E5"))

RESULTADOS = ["Le interesa", "Quiere más información", "No le interesa", "No contestó",
              "Número equivocado", "Ya cerró"]
# El INEGI captura sin acentos; se corrigen las palabras del oficio que más se repiten.
ACENTOS = {"lavanderia": "lavandería", "lavanderias": "lavanderías", "tintoreria": "tintorería",
           "tintorerias": "tintorerías", "planchaduria": "planchaduría", "plachaduria": "planchaduría",
           "planchadurias": "planchadurías"}
PALABRAS_MENORES = {"y", "e", "de", "del", "la", "las", "el", "los", "en", "con", "por", "a"}


def nombre_legible(texto):
    palabras = []
    for i, p in enumerate(texto.strip().split()):
        minuscula = ACENTOS.get(p.lower(), p.lower())
        p = minuscula if minuscula != p.lower() else p
        if i > 0 and minuscula in PALABRAS_MENORES:
            palabras.append(minuscula)
        elif not re.search(r"[aeiou]", minuscula) or re.fullmatch(r"[ivxl]{2,}", minuscula):
            palabras.append(p.upper())  # siglas y números romanos: MR, SL, III
        else:
            palabras.append(p[:1].upper() + p[1:].lower())
    return " ".join(palabras)


def direccion_completa(fila):
    partes = [direccion(fila), colonia(fila)]
    return ", ".join(nombre_legible(p) for p in partes if p)


def contacto_extra(fila):
    return limpio(fila["correoelec"]).lower() or limpio(fila["www"]).lower()


def encabezado(celda):
    celda.font = Font(name=FUENTE, bold=True, color="FFFFFF", size=11)
    celda.fill = PatternFill("solid", fgColor=AZUL)
    celda.alignment = Alignment(wrap_text=True, vertical="center")


def hoja_lista(wb, nombre, filas, con_telefono, tipos):
    hoja = wb.create_sheet(nombre)
    columnas = [("Zona", 22), ("Municipio", 18), ("Lavandería", 36), ("Tipo de lavandería", 20)]
    if con_telefono:
        columnas.append(("Teléfono", 15))
    columnas += [("Cómo acercarse", 46), ("Dirección", 52), ("Correo o página", 30), ("Ubicación", 12),
                 ("¿Ya se le contactó?", 13), ("Resultado", 22), ("Notas", 44)]
    for c, (texto, ancho) in enumerate(columnas, start=1):
        encabezado(hoja.cell(row=1, column=c, value=texto))
        hoja.column_dimensions[get_column_letter(c)].width = ancho
    hoja.row_dimensions[1].height = 30
    indice = {texto: c for c, (texto, _) in enumerate(columnas, start=1)}

    for r, fila in enumerate(filas, start=2):
        valores = {
            "Zona": plaza_de(fila),
            "Municipio": fila["municipio"].strip(),
            "Lavandería": nombre_legible(fila["nom_estab"]),
            "Tipo de lavandería": tipos[fila["id"]],
            "Teléfono": telefono(fila["telefono"]),
            "Dirección": direccion_completa(fila),
            "Correo o página": contacto_extra(fila),
        }
        for texto, c in indice.items():
            celda = hoja.cell(row=r, column=c, value=valores.get(texto))
            celda.font = Font(name=FUENTE, size=10, bold=(texto in {"Lavandería", "Teléfono"}))
            celda.border = LINEA
            celda.alignment = Alignment(vertical="center", wrap_text=(texto in {"Dirección", "Cómo acercarse"}))
        # El consejo sigue al tipo: si quien llama cambia el tipo (por ejemplo, a "Nueva"), el consejo cambia solo.
        consejo = hoja.cell(row=r, column=indice["Cómo acercarse"])
        consejo.value = "=IFERROR(INDEX('Tipos de lavandería'!$C$4:$C$8,MATCH(%s%d,'Tipos de lavandería'!$A$4:$A$8,0)),\"\")" % (get_column_letter(indice["Tipo de lavandería"]), r)
        consejo.font = Font(name=FUENTE, size=9, color=GRIS)
        mapa = hoja.cell(row=r, column=indice["Ubicación"], value="Ver mapa")
        mapa.hyperlink = f"https://www.google.com/maps/search/?api=1&query={fila['latitud']},{fila['longitud']}"
        mapa.font = Font(name=FUENTE, size=10, color=AZUL, underline="single")
        for texto in ("¿Ya se le contactó?", "Resultado", "Notas"):
            hoja.cell(row=r, column=indice[texto]).fill = PatternFill("solid", fgColor=AMARILLO)

    ultima = max(hoja.max_row, 2)
    hoja.freeze_panes = "D2"
    hoja.auto_filter.ref = f"A1:{get_column_letter(len(columnas))}{ultima}"
    si_no = DataValidation(type="list", formula1='"Sí,No"', allow_blank=True)
    resultado = DataValidation(type="list", formula1='"' + ",".join(RESULTADOS) + '"', allow_blank=True)
    hoja.add_data_validation(si_no)
    hoja.add_data_validation(resultado)
    tipo_dv = DataValidation(type="list", formula1='"' + ",".join(TIPOS) + '"', allow_blank=True)
    hoja.add_data_validation(tipo_dv)
    col = lambda texto: get_column_letter(indice[texto])
    tipo_dv.add(f"{col('Tipo de lavandería')}2:{col('Tipo de lavandería')}{ultima}")
    si_no.add(f"{col('¿Ya se le contactó?')}2:{col('¿Ya se le contactó?')}{ultima}")
    resultado.add(f"{col('Resultado')}2:{col('Resultado')}{ultima}")
    rangos = {texto: f"'{nombre}'!${col(texto)}$2:${col(texto)}${ultima}"
              for texto in ("Zona", "¿Ya se le contactó?", "Resultado", "Tipo de lavandería")}
    return rangos


def hoja_resumen(hoja, llamar, visitar, ciudades_fuertes):
    hoja.title = "Resumen"
    hoja.sheet_view.showGridLines = False
    anchos = [30, 16, 16, 16, 16]
    for c, ancho in enumerate(anchos, start=1):
        hoja.column_dimensions[get_column_letter(c)].width = ancho

    hoja["A1"] = "Lavanderías para invitar a Lavandi"
    hoja["A1"].font = Font(name=FUENTE, bold=True, size=20, color=AZUL)
    hoja["A2"] = "Estado de Durango, Torreón y Matamoros  ·  Preparado por el equipo de Nazmelia"
    hoja["A2"].font = Font(name=FUENTE, size=11, color=GRIS)

    # Tres cifras grandes.
    cifras = [
        (f"=COUNTA({llamar['Zona']})+COUNTA({visitar['Zona']})", "lavanderías en la zona"),
        (f"=COUNTA({llamar['Zona']})", "con teléfono, listas para llamar"),
        (ciudades_fuertes, "están en Durango capital y Torreón"),
    ]
    for i, (valor, leyenda) in enumerate(cifras):
        columna = get_column_letter(1 + i * 2) if i else "A"
        celda = hoja[f"{columna}4"]
        celda.value = valor
        celda.font = Font(name=FUENTE, bold=True, size=26, color=AZUL)
        celda.alignment = Alignment(horizontal="left")
        hoja[f"{columna}5"] = leyenda
        hoja[f"{columna}5"].font = Font(name=FUENTE, size=10, color=GRIS)

    hoja["A7"] = "Por zona"
    hoja["A7"].font = Font(name=FUENTE, bold=True, size=13, color=AZUL)
    titulos = ["Zona", "Lavanderías", "Para llamar", "Contactadas", "Interesadas"]
    for c, texto in enumerate(titulos, start=1):
        encabezado(hoja.cell(row=8, column=c, value=texto))
    for k, plaza in enumerate(PLAZAS):
        r = 9 + k
        z = f"$A{r}"
        formulas = [
            plaza,
            f"=COUNTIF({llamar['Zona']},{z})+COUNTIF({visitar['Zona']},{z})",
            f"=COUNTIF({llamar['Zona']},{z})",
            f'=COUNTIFS({llamar["Zona"]},{z},{llamar["¿Ya se le contactó?"]},"Sí")'
            f'+COUNTIFS({visitar["Zona"]},{z},{visitar["¿Ya se le contactó?"]},"Sí")',
            f'=COUNTIFS({llamar["Zona"]},{z},{llamar["Resultado"]},"Le interesa")'
            f'+COUNTIFS({visitar["Zona"]},{z},{visitar["Resultado"]},"Le interesa")',
        ]
        for c, valor in enumerate(formulas, start=1):
            celda = hoja.cell(row=r, column=c, value=valor)
            celda.font = Font(name=FUENTE, size=11)
            celda.border = LINEA
    total_r = 9 + len(PLAZAS)
    hoja.cell(row=total_r, column=1, value="Total")
    for c in range(2, 6):
        letra = get_column_letter(c)
        hoja.cell(row=total_r, column=c, value=f"=SUM({letra}9:{letra}{total_r - 1})")
    for c in range(1, 6):
        celda = hoja.cell(row=total_r, column=c)
        celda.font = Font(name=FUENTE, bold=True, size=11)
        celda.fill = PatternFill("solid", fgColor=AZUL_CLARO)

    r = total_r + 2
    hoja.cell(row=r, column=1, value="Por tipo de lavandería").font = Font(name=FUENTE, bold=True, size=13, color=AZUL)
    for c, texto in enumerate(["Tipo", "Lavanderías", "En Durango capital", "Contactadas", "Interesadas"], start=1):
        encabezado(hoja.cell(row=r + 1, column=c, value=texto))
    for k, tipo in enumerate(TIPOS):
        f = r + 2 + k
        t = f"$A{f}"
        suma = lambda extra: "+".join(f"COUNTIFS({h['Tipo de lavandería']},{t}{extra(h)})" for h in (llamar, visitar))
        valores = [tipo, "=" + suma(lambda h: ""), "=" + suma(lambda h: f',{h["Zona"]},"Durango capital"'),
                   "=" + suma(lambda h: f',{h["¿Ya se le contactó?"]},"Sí"'), "=" + suma(lambda h: f',{h["Resultado"]},"Le interesa"')]
        for c, v in enumerate(valores, start=1):
            celda = hoja.cell(row=f, column=c, value=v); celda.font = Font(name=FUENTE, size=11); celda.border = LINEA
    nota_t = hoja.cell(row=r + 2 + len(TIPOS), column=1, value="«Nueva» no sale de los datos: se marca a mano cuando la lavandería dice que abrió hace poco.")
    nota_t.font = Font(name=FUENTE, italic=True, size=9, color=GRIS)
    r = r + 4 + len(TIPOS)
    hoja.cell(row=r, column=1, value="Cómo usarlo").font = Font(name=FUENTE, bold=True, size=13, color=AZUL)
    pasos = [
        "1.  Empiece por la hoja «Para llamar»: cada renglón es una lavandería con su teléfono.",
        "2.  Después de cada llamada, llene las tres columnas amarillas: si ya se le contactó, el resultado y una nota.",
        "3.  La hoja «Para visitar» trae las que no tienen teléfono: con «Ver mapa» se llega a la puerta.",
        "4.  La columna «Tipo de lavandería» dice cómo acercarse a cada una. Si al hablar descubre que es nueva, cámbiela.",
        "5.  Este resumen se actualiza solo con lo que se anote.",
    ]
    for i, texto in enumerate(pasos, start=1):
        celda = hoja.cell(row=r + i, column=1, value=texto)
        celda.font = Font(name=FUENTE, size=11)
    r += len(pasos) + 2
    nota = hoja.cell(row=r, column=1, value="Datos tomados del registro oficial de negocios del INEGI. "
                                            "Algún teléfono o dirección puede haber cambiado.")
    nota.font = Font(name=FUENTE, italic=True, size=9, color=GRIS)


def hoja_tipos(wb):
    hoja = wb.create_sheet("Tipos de lavandería")
    hoja.sheet_view.showGridLines = False
    hoja["A1"] = "Cinco tipos de lavandería, cinco maneras de acercarse"
    hoja["A1"].font = Font(name=FUENTE, bold=True, size=16, color=AZUL)
    hoja["A2"] = "El detalle de cada uno está en la guía «Cómo hablarle a cada lavandería» del kit."
    hoja["A2"].font = Font(name=FUENTE, size=10, color=GRIS)
    for c, (t, w) in enumerate([("Tipo", 24), ("Cómo reconocerla", 50), ("Cómo acercarse", 70)], start=1):
        encabezado(hoja.cell(row=3, column=c, value=t)); hoja.column_dimensions[get_column_letter(c)].width = w
    for k, tipo in enumerate(TIPOS):
        for c, v in enumerate([tipo, *TIPOS_INFO[tipo]], start=1):
            celda = hoja.cell(row=4 + k, column=c, value=v)
            celda.font = Font(name=FUENTE, size=11, bold=(c == 1)); celda.border = LINEA
            celda.alignment = Alignment(wrap_text=True, vertical="top")
        hoja.row_dimensions[4 + k].height = 34


def para_imprimir(hoja):
    hoja.page_setup.orientation = "landscape"
    hoja.page_setup.fitToWidth = 1
    hoja.page_setup.fitToHeight = 0
    hoja.sheet_properties.pageSetUpPr.fitToPage = True
    hoja.print_title_rows = "1:1" if hoja.title != "Resumen" else None


def main():
    carpeta = Path(sys.argv[1]) if len(sys.argv) > 1 else Path.cwd() / "descarga_denue"
    salida = Path(sys.argv[2]) if len(sys.argv) > 2 else SALIDA_POR_OMISION
    carpeta.mkdir(parents=True, exist_ok=True)
    lavanderias, _, _ = registros(carpeta)

    orden = {p: i for i, p in enumerate(PLAZAS)}
    lavanderias.sort(key=lambda f: (orden[plaza_de(f)], normalizar(f["municipio"]), normalizar(f["nom_estab"])))
    con_tel = [f for f in lavanderias if telefono(f["telefono"])]
    sin_tel = [f for f in lavanderias if not telefono(f["telefono"])]

    por_plaza = Counter(plaza_de(f) for f in lavanderias)
    fuertes = por_plaza["Durango capital"] + por_plaza["Torreón"]
    ciudades_fuertes = f"{round(10 * fuertes / len(lavanderias))} de cada 10"

    wb = Workbook()
    resumen = wb.active
    tipos = tipos_de_prospecto(lavanderias)
    llamar = hoja_lista(wb, "Para llamar", con_tel, con_telefono=True, tipos=tipos)
    visitar = hoja_lista(wb, "Para visitar", sin_tel, con_telefono=False, tipos=tipos)
    hoja_tipos(wb)
    hoja_resumen(resumen, llamar, visitar, ciudades_fuertes)
    for hoja in wb.worksheets:
        para_imprimir(hoja)
    wb.properties.creator = "Nazmelia"
    wb.properties.title = "Lavanderías para invitar a Lavandi"
    salida.parent.mkdir(parents=True, exist_ok=True)
    wb.save(salida)
    print(f"Total {len(lavanderias)} · para llamar {len(con_tel)} · para visitar {len(sin_tel)} · {ciudades_fuertes}")
    print(f"Escrito: {salida}")


if __name__ == "__main__":
    main()
