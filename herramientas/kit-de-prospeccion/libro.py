# -*- coding: utf-8 -*-
"""El libro de resultados: donde se pasan las hojas de visita y las entrevistas. El resumen se calcula con fórmulas,
así que se actualiza solo. Mide lo que el negocio necesita saber: cuánto cobran las lavanderías de la zona, cuántas
aceptan y con qué comisión, y si la gente pediría por una aplicación (hipótesis H-001, H-002 y H-006)."""
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

FUENTE = 'Arial'
AZUL, AZUL_CLARO, AMARILLO, GRIS = '297899', 'E3F1F6', 'FFF4CC', '515A64'
LINEA = Border(bottom=Side(style='thin', color='DCE0E5'))
FILAS = 300  # renglones preparados para llenar
TIPOS = ['De barrio', 'Establecida', 'Nueva', 'Cadena', 'Tintorería o planchaduría']

HOJAS = {
    'Visitas': [
        ('Lavandería', 28, None), ('Zona o colonia', 20, None), ('Tipo de lavandería', 20, TIPOS),
        ('¿Lleva menos de un año abierta?', 12, ['Sí', 'No']), ('Fecha', 12, None), ('Quién visitó', 16, None),
        ('Precio por kilo ($)', 12, 'num'), ('Mínimo de kilos', 10, 'num'), ('Planchado ($)', 11, 'num'),
        ('¿Entrega a domicilio?', 16, ['No', 'Sí, sin costo', 'Sí, cobra envío']), ('Costo del envío ($)', 12, 'num'),
        ('Pedidos al día', 10, 'num'), ('¿Le cabrían más?', 12, ['Sí', 'No']), ('Cómo le cobran', 16, None),
        ('Cómo le piden', 16, None), ('¿Quiere ser fundadora?', 18, ['Sí, firmó', 'Lo va a pensar', 'No']),
        ('La comisión le parece', 16, ['Bien', 'Alta', 'No acepta']), ('Comisión que le parece justa', 16, None),
        ('Lo que le preocupa', 36, None), ('Siguiente paso', 16, ['Mandar hoja', 'Volver a visitar', 'Firmó carta', 'Nada más']),
    ],
    'Entrevistas': [
        ('Colonia', 20, None), ('Fecha', 12, None), ('Precio que se probó ($)', 12, 'num'),
        ('¿Tiene lavadora?', 12, ['Sí', 'No']), ('¿Usa lavandería?', 14, ['Cada semana', 'Cada mes', 'A veces', 'Nunca']),
        ('Pagó la última vez ($)', 12, 'num'), ('Kilos', 8, 'num'), ('Cómo la pide', 16, ['En persona', 'WhatsApp', 'Le recogen en casa']),
        ('Cómo paga', 14, ['Efectivo', 'Transferencia', 'Tarjeta']), ('¿Usaría Lavandi a ese precio?', 16, ['Sí, seguro', 'Tal vez', 'No']),
        ('Por qué', 32, None), ('Qué le daría confianza', 32, None), ('Qué le haría dejarla', 32, None), ('¿Quiere aviso?', 10, ['Sí', 'No']),
    ],
    'Lavan en casa': [
        ('Colonia', 20, None), ('Fecha', 12, None), ('¿Ya lava ropa ajena?', 14, ['Sí', 'No, le interesa']), ('Familias que atiende', 10, 'num'),
        ('Cobra ($)', 10, 'num'), ('Por (kilo, carga, prenda)', 14, None), ('Kilos que podría al día', 12, 'num'),
        ('Cómo entrega', 18, None), ('Cómo le pagan', 16, None), ('¿Da recibo o factura?', 12, ['Sí', 'No']),
        ('Lo que le preocupa', 32, None), ('Lo que necesitaría', 32, None), ('¿Se le puede volver a buscar?', 14, ['Sí', 'No']),
    ],
}


def encabezado(celda):
    celda.font = Font(name=FUENTE, bold=True, color='FFFFFF', size=10)
    celda.fill = PatternFill('solid', fgColor=AZUL)
    celda.alignment = Alignment(wrap_text=True, vertical='center')


def hoja_datos(wb, nombre, columnas):
    h = wb.create_sheet(nombre)
    for c, (titulo, ancho, tipo) in enumerate(columnas, 1):
        encabezado(h.cell(row=1, column=c, value=titulo))
        h.column_dimensions[get_column_letter(c)].width = ancho
        letra = get_column_letter(c)
        if isinstance(tipo, list):
            dv = DataValidation(type='list', formula1='"%s"' % ','.join(tipo), allow_blank=True)
            h.add_data_validation(dv); dv.add('%s2:%s%d' % (letra, letra, FILAS + 1))
        elif tipo == 'num':
            dv = DataValidation(type='decimal', operator='greaterThanOrEqual', formula1='0', allow_blank=True,
                                error='Escriba solo el número, sin signo de pesos.', showErrorMessage=True)
            h.add_data_validation(dv); dv.add('%s2:%s%d' % (letra, letra, FILAS + 1))
            for r in range(2, FILAS + 2):
                h['%s%d' % (letra, r)].number_format = '#,##0.##'
    for r in range(2, FILAS + 2):
        for c in range(1, len(columnas) + 1):
            h.cell(row=r, column=c).font = Font(name=FUENTE, size=10)
            h.cell(row=r, column=c).border = LINEA
    h.row_dimensions[1].height = 42
    h.freeze_panes = 'B2'
    h.auto_filter.ref = 'A1:%s%d' % (get_column_letter(len(columnas)), FILAS + 1)
    return {t: "'%s'!$%s$2:$%s$%d" % (nombre, get_column_letter(c), get_column_letter(c), FILAS + 1) for c, (t, _, _) in enumerate(columnas, 1)}


def construir(ruta):
    wb = Workbook()
    res = wb.active; res.title = 'Resumen'
    v = hoja_datos(wb, 'Visitas', HOJAS['Visitas'])
    en = hoja_datos(wb, 'Entrevistas', HOJAS['Entrevistas'])
    lc = hoja_datos(wb, 'Lavan en casa', HOJAS['Lavan en casa'])

    res.sheet_view.showGridLines = False
    res.column_dimensions['A'].width = 58; res.column_dimensions['B'].width = 16; res.column_dimensions['C'].width = 50
    res['A1'] = 'Resultados de la prospección'; res['A1'].font = Font(name=FUENTE, bold=True, size=18, color=AZUL)
    res['A2'] = 'Se calcula solo con lo que se anote en las hojas Visitas, Entrevistas y Lavan en casa.'
    res['A2'].font = Font(name=FUENTE, size=10, color=GRIS)

    pct = lambda num, den: '=IFERROR(%s/%s,"-")' % (num, den)
    nv = 'COUNTA(%s)' % v['Lavandería']
    ne = 'COUNTA(%s)' % en['Colonia']
    nl = 'COUNTA(%s)' % lc['Colonia']
    bloques = [
        ('Lavanderías', [
            ('Lavanderías visitadas', '=' + nv, '0', ''),
            ('Quieren ser fundadoras y firmaron la carta', '=COUNTIF(%s,"Sí, firmó")' % v['¿Quiere ser fundadora?'], '0', ''),
            ('… de cada 100 visitadas', pct('COUNTIF(%s,"Sí, firmó")' % v['¿Quiere ser fundadora?'], nv), '0%', 'Cuántas dicen que sí de cada 100 que se visitan.'),
            ('Les parece bien la comisión', pct('COUNTIF(%s,"Bien")' % v['La comisión le parece'], 'COUNTA(%s)' % v['La comisión le parece']), '0%', 'Entre las que respondieron a la pregunta de la comisión.'),
            ('Precio por kilo, promedio ($)', '=IFERROR(AVERAGE(%s),"-")' % v['Precio por kilo ($)'], '#,##0.00', ''),
            ('Precio por kilo, el más bajo ($)', '=IF(COUNT(%s)=0,"-",MIN(%s))' % (v['Precio por kilo ($)'], v['Precio por kilo ($)']), '#,##0.00', ''),
            ('Precio por kilo, el más alto ($)', '=IF(COUNT(%s)=0,"-",MAX(%s))' % (v['Precio por kilo ($)'], v['Precio por kilo ($)']), '#,##0.00', ''),
            ('Ya entregan a domicilio', pct('COUNTIF(%s,"Sí*")' % v['¿Entrega a domicilio?'], 'COUNTA(%s)' % v['¿Entrega a domicilio?']), '0%', ''),
            ('… y cobran el envío', pct('COUNTIF(%s,"Sí, cobra envío")' % v['¿Entrega a domicilio?'], 'COUNTIF(%s,"Sí*")' % v['¿Entrega a domicilio?']), '0%', 'De las que entregan a domicilio.'),
            ('Les cabrían más pedidos', pct('COUNTIF(%s,"Sí")' % v['¿Le cabrían más?'], 'COUNTA(%s)' % v['¿Le cabrían más?']), '0%', ''),
        ]),
        ('Por tipo de lavandería: visitadas y las que firmaron', [
            ('%s: visitadas' % tipo, '=COUNTIF(%s,"%s")' % (v['Tipo de lavandería'], tipo), '0', '') for tipo in TIPOS
        ] + [
            ('%s: firmaron la carta' % tipo, '=COUNTIFS(%s,"%s",%s,"Sí, firmó")' % (v['Tipo de lavandería'], tipo, v['¿Quiere ser fundadora?']), '0',
             'Dice con qué tipo de lavandería conviene insistir.' if tipo == TIPOS[0] else '') for tipo in TIPOS
        ]),
        ('Personas entrevistadas', [
            ('Personas entrevistadas', '=' + ne, '0', 'Con 30 ya se ve si la idea se sostiene.'),
            ('Tienen lavadora en casa', pct('COUNTIF(%s,"Sí")' % en['¿Tiene lavadora?'], 'COUNTA(%s)' % en['¿Tiene lavadora?']), '0%', ''),
            ('Usarían Lavandi al precio que se probó: sí, seguro', pct('COUNTIF(%s,"Sí, seguro")' % en['¿Usaría Lavandi a ese precio?'], 'COUNTA(%s)' % en['¿Usaría Lavandi a ese precio?']), '0%', ''),
            ('… tal vez', pct('COUNTIF(%s,"Tal vez")' % en['¿Usaría Lavandi a ese precio?'], 'COUNTA(%s)' % en['¿Usaría Lavandi a ese precio?']), '0%', ''),
            ('Pagan en efectivo', pct('COUNTIF(%s,"Efectivo")' % en['Cómo paga'], 'COUNTA(%s)' % en['Cómo paga']), '0%', ''),
            ('Lo que pagaron la última vez, promedio ($)', '=IFERROR(AVERAGE(%s),"-")' % en['Pagó la última vez ($)'], '#,##0.00', ''),
            ('Quieren que les avisen cuando abra', '=COUNTIF(%s,"Sí")' % en['¿Quiere aviso?'], '0', ''),
        ]),
        ('Personas que lavan en casa', [
            ('Conversaciones', '=' + nl, '0', 'Solo para explorar. Todavía no se les puede sumar.'),
            ('Ya lavan ropa ajena', pct('COUNTIF(%s,"Sí")' % lc['¿Ya lava ropa ajena?'], 'COUNTA(%s)' % lc['¿Ya lava ropa ajena?']), '0%', ''),
            ('Kilos que podrían atender al día, promedio', '=IFERROR(AVERAGE(%s),"-")' % lc['Kilos que podría al día'], '#,##0.0', ''),
            ('Se les puede volver a buscar', '=COUNTIF(%s,"Sí")' % lc['¿Se le puede volver a buscar?'], '0', ''),
        ]),
    ]
    r = 4
    for titulo, filas in bloques:
        res.cell(row=r, column=1, value=titulo).font = Font(name=FUENTE, bold=True, size=13, color=AZUL)
        r += 1
        for etiqueta, formula, formato, nota in filas:
            a = res.cell(row=r, column=1, value=etiqueta); a.font = Font(name=FUENTE, size=11); a.border = LINEA
            b = res.cell(row=r, column=2, value=formula); b.font = Font(name=FUENTE, size=12, bold=True); b.number_format = formato
            b.border = LINEA; b.alignment = Alignment(horizontal='right')
            c = res.cell(row=r, column=3, value=nota); c.font = Font(name=FUENTE, size=9, color=GRIS); c.border = LINEA
            r += 1
        r += 1
    res.cell(row=r, column=1, value='Cómo se llena').font = Font(name=FUENTE, bold=True, size=13, color=AZUL)
    pasos = [
        'Cada hoja de visita en papel es un renglón de la hoja Visitas; cada entrevista, un renglón de Entrevistas.',
        'Las columnas con flecha se eligen de una lista. Los montos van como número, sin signo de pesos: 28, no $28.',
        'Ejemplo de un renglón de Visitas: Lavandería La Burbuja · Centro · Precio por kilo 28 · ¿Entrega a domicilio? Sí, cobra envío · Costo del envío 30 · ¿Quiere ser fundadora? Lo va a pensar.',
        'Las llamadas se siguen anotando en la lista de lavanderías. Este libro es para las visitas y las entrevistas.',
    ]
    for i, t in enumerate(pasos, 1):
        c = res.cell(row=r + i, column=1, value=t); c.font = Font(name=FUENTE, size=10); c.alignment = Alignment(wrap_text=True)
        res.merge_cells(start_row=r + i, start_column=1, end_row=r + i, end_column=3)
        res.row_dimensions[r + i].height = 28
    wb.properties.creator = 'Lavandi'; wb.properties.title = 'Resultados de la prospección'
    wb.save(ruta)
